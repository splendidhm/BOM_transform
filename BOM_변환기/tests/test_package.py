"""Integration tests exercise Qt drop events, GUI conversion, and saved XLSX files."""
import os,sys,time,json,tempfile,shutil,csv
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'vendor'))
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from PyQt5.QtWidgets import QApplication
from PyQt5.QtCore import QMimeData,QUrl,Qt,QPoint,QPointF
from PyQt5.QtGui import QDragEnterEvent,QDropEvent
from app import MainWindow,BatchWorker
from bom_core import read_raw,map_rows,convert_one,validate_template,BomError
import openpyxl
from copy import copy
from decimal import Decimal

SOURCE=Path(sys.argv[1]);TEMPLATE=ROOT/'templates/BOM 자동 입력 양식.xlsx'
OUT=ROOT/'완료파일';REPORT=ROOT/'tests/test_results.json'
results=[]
def check(name,fn):
    fn();results.append({'test':name,'result':'PASS'});print('PASS',name,flush=True)
def wait_until(fn,seconds=120):
    end=time.monotonic()+seconds
    while not fn():
        if time.monotonic()>end:raise AssertionError('Timed out')
        app.processEvents();time.sleep(.015)
    app.processEvents()
app=QApplication.instance() or QApplication([])
window=MainWindow();window.show();app.processEvents()
window.template.setText(str(TEMPLATE));window.output.setText(str(OUT));window.overwrite.setChecked(True)

def drop_test():
    mime=QMimeData();mime.setUrls([QUrl.fromLocalFile(str(SOURCE/(c+'.xls'))) for c in ['110585','150149']])
    viewport=window.table.viewport()
    enter=QDragEnterEvent(QPoint(20,20),Qt.CopyAction,mime,Qt.LeftButton,Qt.NoModifier)
    QApplication.sendEvent(viewport,enter);assert enter.isAccepted()
    event=QDropEvent(QPointF(20,20),Qt.CopyAction,mime,Qt.LeftButton,Qt.NoModifier)
    QApplication.sendEvent(viewport,event)
    assert event.isAccepted() and len(window.paths)==2
    window.add_paths([str(SOURCE/'150149.xls')]);assert len(window.paths)==2
check('Qt file drop and duplicate-path suppression',drop_test)

def gui_convert():
    window.convert_button.click();assert window.busy
    wait_until(lambda:not window.busy)
    assert window.last_summary['totals']['완료']==2,window.last_summary
    assert window.progress.value()==2
    assert Path(window.last_summary['log']).is_file()
    for i in range(2):assert window.table.item(i,2).text()=='완료'
check('GUI button sequentially converts both supplied RAW files',gui_convert)

def independent_verify():
    template=openpyxl.load_workbook(TEMPLATE).active
    explicit_cells=set(template._cells)
    for code in ['110585','150149']:
        raw=list(csv.reader((SOURCE/(code+'.xls')).read_text(encoding='cp949').splitlines(),delimiter='\t'))[1:]
        ordered=sorted(enumerate(raw,2),key=lambda p:(int(next(v for v in p[1][3:18] if v.strip())),p[0]))
        from collections import Counter
        children=Counter(r[18] for r in raw)
        wb=openpyxl.load_workbook(OUT/(code+' 완료.xlsx'));sheet=wb.active
        assert wb.sheetnames==[code] and sheet.max_row==len(raw)+4
        for index,(_,r) in enumerate(ordered,5):
            want=[int(r[18]),'KG','M',0,1,int(r[22]),r[25],Decimal(r[26]),children[r[22]]]+[None]*11
            for col,value in enumerate(want,1):
                actual=sheet.cell(index,col).value
                if col==8:actual=Decimal(str(actual))
                assert actual==value,(code,index,col)
        for row in sheet:
            for c in row:
                if (c.row,c.column) in explicit_cells:
                    ref=template.cell(c.row,c.column)
                    if c.row<5:assert c.value==ref.value
                    for attr in ['font','fill','border','alignment','number_format','protection']:
                        assert copy(getattr(c,attr))==copy(getattr(ref,attr)),(code,c.coordinate,attr)
                elif c.value is not None:
                    # Newly populated cells without explicit formatting inherit cellXf 0.
                    assert c._style==template.parent._cell_styles[0],(code,c.coordinate,'default style')
check('Independent 51-row / 1020-data-cell RAW verification and template styles',independent_verify)

def skip_test():
    before=(OUT/'150149 완료.xlsx').read_bytes()
    result=convert_one(SOURCE/'150149.xls',TEMPLATE,OUT,False)
    assert result['status']=='건너뜀' and before==(OUT/'150149 완료.xlsx').read_bytes()
check('Existing-file protection',skip_test)

with tempfile.TemporaryDirectory(prefix='bom-gui-tests-',dir=ROOT/'tests') as temp:
    temp=Path(temp)
    def large_queue():
        folder=temp/'queue';folder.mkdir()
        for i in range(1200):(folder/(str(700000+i)+'.xls')).write_bytes(b'')
        other=MainWindow();other.add_paths([str(folder)]);assert len(other.paths)==1200
        other.clear_files();assert not other.paths;other.close()
    check('1200-file folder queue',large_queue)
    def expanded_rows():
        folder=temp/'large';folder.mkdir();raw=folder/'150149.xls'
        lines=(SOURCE/'150149.xls').read_text(encoding='cp949').splitlines()
        raw.write_text('\n'.join([lines[0]]+lines[1:]*3),encoding='cp949')
        result=convert_one(raw,TEMPLATE,temp/'expanded')
        assert result['rows']==57
        sheet=openpyxl.load_workbook(result['output']).active
        assert sheet.max_row==61 and sheet['G61'].value is not None
    check('57-row BOM expands beyond the 32-row blank template',expanded_rows)
    def invalid_continue():
        bad=temp/'999999.xls';bad.write_bytes((SOURCE/'150149.xls').read_bytes())
        worker=BatchWorker([bad,SOURCE/'150149.xls'],TEMPLATE,temp/'continue');done=[]
        worker.completed.connect(done.append);worker.start();wait_until(lambda:not worker.isRunning());app.processEvents()
        assert done[0]['totals']['실패']==1 and done[0]['totals']['완료']==1
    check('Bad input fails independently and next file completes',invalid_continue)
    def duplicate_code():
        folder=temp/'duplicate';folder.mkdir();duplicate=folder/'150149.xls';duplicate.write_bytes((SOURCE/'150149.xls').read_bytes())
        worker=BatchWorker([SOURCE/'150149.xls',duplicate],TEMPLATE,OUT,True);done=[]
        worker.completed.connect(done.append);worker.start();wait_until(lambda:not worker.isRunning());app.processEvents()
        assert done[0]['totals']['완료']==1 and done[0]['totals']['건너뜀']==1
    check('Same product in two directories does not overwrite within a batch',duplicate_code)
    def cancellation():
        worker=BatchWorker([SOURCE/'110585.xls',SOURCE/'150149.xls'],TEMPLATE,temp/'cancel');done=[]
        worker.started_file.connect(lambda index:worker.requestInterruption())
        worker.completed.connect(done.append);worker.start();wait_until(lambda:not worker.isRunning());app.processEvents()
        assert done[0]['totals']['완료']==1 and done[0]['totals']['취소']==1,done
        assert not (temp/'cancel/150149 완료.xlsx').exists()
    check('Cancellation finishes current file and cancels remaining queue',cancellation)
    def unit_and_validation():
        product,rows=read_raw(SOURCE/'150149.xls');assert product=='150149'
        unit_folder=temp/'units';unit_folder.mkdir();filename=unit_folder/'150149.xls'
        table=list(csv.reader((SOURCE/'150149.xls').read_text(encoding='cp949').splitlines(),delimiter='\t'))
        table[1][25]=' kg '
        with filename.open('w',encoding='cp949',newline='') as f:csv.writer(f,delimiter='\t').writerows(table)
        _,values=map_rows(read_raw(filename)[1]);assert values[0][6]==' kg '
        table[1][26]='NaN'
        with filename.open('w',encoding='cp949',newline='') as f:csv.writer(f,delimiter='\t').writerows(table)
        try:read_raw(filename)
        except BomError:pass
        else:raise AssertionError('NaN accepted')
    check('Unit text is exact, non-finite quantity is rejected',unit_and_validation)

window.status.setText('검증 완료: 2개 제품, 51개 BOM 관계 매핑 통과')
window.grab().save(str(ROOT/'tests/gui_preview.png'))
window.close()
REPORT.write_text(json.dumps({'passed':True,'tests':results},ensure_ascii=False,indent=2),encoding='utf-8')
print('ALL TESTS PASSED',flush=True)

"""BOM batch engine: read RAW, author via Artifact Tool, preserve template, verify."""
from __future__ import annotations
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from copy import deepcopy, copy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parent
RUNTIME=Path(os.environ.get('BOM_RUNTIME',str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies')))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q=lambda name:'{'+NS+'}'+name
EXPECTED=['ITEM_CD','M_UNIT_CD','PRPOS_CD','ALT_NO','BOM_QT','MTRIL_CD','UNIT_CD','NET_QT','SCRAP_RT']

class BomError(ValueError): pass

@dataclass(frozen=True)
class RawRow:
    source_row:int
    level:int
    parent:str
    child:str
    unit:str
    quantity:Decimal

def cell_text(value):
    if value is None: return ''
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value)

def read_raw(raw:Path):
    raw=Path(raw)
    product=raw.stem
    if not re.fullmatch(r'[0-9]+',product): raise BomError('파일명은 제품코드 숫자여야 합니다.')
    payload=raw.read_bytes()
    if payload.startswith(b'PK'):
        # BytesIO also supports XLSX data carrying an .xls suffix.
        book=openpyxl.load_workbook(io.BytesIO(payload),read_only=True,data_only=False)
        try: table=[list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
        finally: book.close()
    elif payload.startswith(b'\xd0\xcf\x11\xe0'):
        import xlrd
        book=xlrd.open_workbook(file_contents=payload)
        try:
            sheet=book.sheet_by_index(0)
            table=[sheet.row_values(i) for i in range(sheet.nrows)]
        finally: book.release_resources()
    else:
        text=None
        for encoding in ['utf-8-sig','cp949']:
            try: text=payload.decode(encoding);break
            except UnicodeDecodeError: continue
        if text is None: raise BomError('UTF-8 또는 CP949 텍스트가 아닙니다.')
        table=list(csv.reader(io.StringIO(text),delimiter='\t'))
    if not table: raise BomError('빈 RAW 파일입니다.')
    headers=[cell_text(v).strip() for v in table[0]]
    for col,title in [(0,'제품코드'),(18,'모품목코드'),(22,'자품목코드'),(25,'단위'),(26,'소요량')]:
        if len(headers)<=col or headers[col]!=title: raise BomError(f'RAW {col+1}열 헤더가 {title}와 다릅니다.')
    result=[]
    for i,values in enumerate(table[1:],2):
        cells=[cell_text(v) for v in values]
        if not any(x.strip() for x in cells): continue
        if len(cells)<27: raise BomError(f'{i}행: 필수 열이 부족합니다.')
        if cells[0].strip()!=product: raise BomError(f'{i}행: 파일명과 제품코드가 다릅니다.')
        levels=[(j+1,v.strip()) for j,v in enumerate(cells[3:18]) if v.strip()]
        if len(levels)!=1 or levels[0][1]!=str(levels[0][0]): raise BomError(f'{i}행: BOM 단계가 잘못되었습니다.')
        parent,child=cells[18].strip(),cells[22].strip()
        if not re.fullmatch(r'[0-9]+',parent) or not re.fullmatch(r'[0-9]+',child): raise BomError(f'{i}행: 모품목/자품목코드를 확인하세요.')
        if not cells[25].strip(): raise BomError(f'{i}행: 단위가 비어 있습니다.')
        try: qty=Decimal(cells[26].strip())
        except InvalidOperation: raise BomError(f'{i}행: 소요량이 숫자가 아닙니다.')
        if not qty.is_finite() or qty<0 or qty.as_tuple().exponent < -6:
            raise BomError(f'{i}행: 소요량은 0 이상, 소수점 6자리 이내여야 합니다.')
        if qty>=Decimal('10000000000000'): raise BomError(f'{i}행: 소요량이 양식의 숫자 길이를 초과합니다.')
        result.append(RawRow(i,levels[0][0],parent,child,cells[25],qty))
    if not result: raise BomError('변환할 BOM 데이터가 없습니다.')
    return product,result

def map_rows(rows):
    counts=Counter(r.parent for r in rows)
    def code(v): return v if v.startswith('0') or int(v)>9007199254740991 else int(v)
    ordered=sorted(rows,key=lambda r:(r.level,r.source_row))
    values=[[code(r.parent),'KG','M',0,1,code(r.child),r.unit,float(r.quantity),counts[r.child]]+[None]*11 for r in ordered]
    return ordered,values

def validate_template(template):
    wb=openpyxl.load_workbook(template,read_only=True,data_only=False)
    try:
        if len(wb.sheetnames)!=1: raise BomError('양식은 시트가 하나인 파일을 선택하세요.')
        sheet=wb.worksheets[0]
        if [sheet.cell(1,i).value for i in range(1,10)]!=EXPECTED: raise BomError('BOM 양식의 A~I 헤더가 다릅니다.')
    finally: wb.close()

def preserve_template(template:Path,stage:Path,product:str,count:int):
    """Keep every original ZIP part; extend row 5 styling when the RAW is larger."""
    with zipfile.ZipFile(template) as src,zipfile.ZipFile(stage) as built:
        sheet_path='xl/worksheets/sheet1.xml'
        original=ET.fromstring(src.read(sheet_path))
        authored=ET.fromstring(built.read(sheet_path))
        data=original.find(Q('sheetData'))
        row5=next((r for r in data if r.get('r')=='5'),None)
        prototype=deepcopy(row5) if row5 is not None else ET.Element(Q('row'),{'r':'5'})
        existing={int(r.get('r')):deepcopy(r) for r in data if int(r.get('r'))>=5}
        for row in list(data):
            if int(row.get('r'))>=5: data.remove(row)
        cells={}
        for index in range(5,count+5):
            row=existing.get(index,deepcopy(prototype));row.set('r',str(index))
            for c in list(row):
                if c.tag!=Q('c'): continue
                col=re.sub(r'[0-9]+$','',c.get('r'))
                c.set('r',col+str(index));c.attrib.pop('t',None)
                for value in list(c):
                    if value.tag in [Q('v'),Q('f'),Q('is')]: c.remove(value)
                cells[c.get('r')]=c
            data.append(row)
        by_row={int(r.get('r')):r for r in data}
        strings=[]
        if 'xl/sharedStrings.xml' in built.namelist():
            strings=[''.join(si.itertext()) for si in ET.fromstring(built.read('xl/sharedStrings.xml'))]
        for c in authored.iter(Q('c')):
            address=c.get('r');index=int(re.search(r'[0-9]+$',address)[0])
            if index<5 or index>count+4: continue
            if not any(c.find(Q(k)) is not None for k in ['v','f','is']): continue
            dest=cells.get(address)
            if dest is None:
                dest=ET.SubElement(by_row[index],Q('c'),{'r':address});cells[address]=dest
            if c.get('t')=='s':
                dest.set('t','inlineStr');node=ET.SubElement(ET.SubElement(dest,Q('is')),Q('t'))
                node.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                node.text=strings[int(c.find(Q('v')).text)]
            else:
                if c.get('t'):dest.set('t',c.get('t'))
                for value in c:
                    if value.tag in [Q('v'),Q('f'),Q('is')]:dest.append(deepcopy(value))
        # Correct cell order if the blank template omitted any data cells.
        def col_key(c):
            letters=re.sub(r'[0-9]+$','',c.get('r','A1'));n=0
            for letter in letters:n=n*26+ord(letter)-64
            return n
        for row in data:
            if int(row.get('r'))>=5:row[:]=sorted(row,key=col_key)
        dimension=original.find(Q('dimension'))
        if dimension is not None:dimension.set('ref',f'A1:T{count+4}')
        book=ET.fromstring(src.read('xl/workbook.xml'))
        sheet=book.find(Q('sheets'))[0];old=sheet.get('name');sheet.set('name',product)
        for defined in book.iter(Q('definedName')):
            if defined.text:defined.text=defined.text.replace("'"+old+"'!","'"+product+"'!").replace(old+'!',product+'!')
        # lxml preserves the existing namespace declarations needed by Office extensions.
        from lxml import etree as LE
        def retain_namespaces(original_bytes,new_root):
            old_root=LE.fromstring(original_bytes)
            new_root=LE.fromstring(ET.tostring(new_root,encoding='utf-8'))
            rebuilt=LE.Element(old_root.tag,attrib=dict(new_root.attrib),nsmap=old_root.nsmap)
            rebuilt.text=new_root.text
            for child in new_root:rebuilt.append(child)
            return LE.tostring(rebuilt,encoding='utf-8',xml_declaration=True)
        buffer=io.BytesIO()
        with zipfile.ZipFile(buffer,'w') as final:
            for item in src.infolist():
                payload=src.read(item.filename)
                if item.filename==sheet_path:payload=retain_namespaces(payload,original)
                elif item.filename=='xl/workbook.xml':payload=retain_namespaces(payload,book)
                final.writestr(item,payload)
    stage.write_bytes(buffer.getvalue())

def verify_output(output,template,product,ordered,expected):
    wb=openpyxl.load_workbook(output);base=openpyxl.load_workbook(template)
    try:
        sheet=wb.active;ref=base.active
        if wb.sheetnames!=[product] or sheet.max_row!=len(ordered)+4:raise BomError('출력 시트명/행 수 검증 실패')
        for row in range(1,5):
            for col in range(1,21):
                if sheet.cell(row,col).value!=ref.cell(row,col).value:raise BomError('양식 안내행 검증 실패')
        for row,(raw,values) in enumerate(zip(ordered,expected),5):
            for col,want in enumerate(values,1):
                got=sheet.cell(row,col).value
                if col==8:
                    if Decimal(str(got))!=raw.quantity:raise BomError(f'{row}행: 수량 저장 검증 실패')
                elif got!=want:raise BomError(f'{row}행 {col}열: 저장 검증 실패')
        return len(ordered)
    finally:wb.close();base.close()

def convert_one(raw,template,out_dir,overwrite=False):
    raw,template,out_dir=Path(raw).resolve(),Path(template).resolve(),Path(out_dir).resolve()
    product,rows=read_raw(raw)
    target=out_dir/(product+' 완료.xlsx')
    if target in [raw,template]:raise BomError('출력 파일이 입력 파일과 같습니다.')
    if target.exists() and not overwrite:return {'status':'건너뜀','rows':0,'output':str(target),'message':'같은 이름의 완료 파일이 있습니다.'}
    validate_template(template)
    ordered,values=map_rows(rows)
    node=RUNTIME/'node/bin/node.exe'
    if not node.is_file():raise BomError('Node 실행 환경을 찾을 수 없습니다. 사용안내의 실행 환경을 확인하세요.')
    out_dir.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.bom-',dir=out_dir) as tmp:
        tmp=Path(tmp);stage=tmp/'stage.xlsx';job=tmp/'job.json'
        job.write_text(json.dumps({'template':str(template),'stage':str(stage),'values':values},ensure_ascii=False),encoding='utf-8')
        env=dict(os.environ)
        env['ARTIFACT_TOOL_ENTRY']=str(RUNTIME/'node/node_modules/@oai/artifact-tool/dist/artifact_tool.mjs')
        run=subprocess.run([str(node),str(ROOT/'artifact_engine.mjs'),str(job)],capture_output=True,encoding='utf-8',errors='replace',
            timeout=300,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0),env=env)
        if run.returncode:raise BomError('엑셀 생성 실패: '+(run.stderr or run.stdout)[-900:])
        preserve_template(template,stage,product,len(rows))
        verify_output(stage,template,product,ordered,values)
        # No partial files or replacement of existing files unless explicitly selected.
        if overwrite:os.replace(stage,target)
        else:
            try:os.rename(stage,target)
            except FileExistsError:return {'status':'건너뜀','rows':0,'output':str(target),'message':'처리 중 같은 이름의 파일이 생성되었습니다.'}
    return {'status':'완료','rows':len(rows),'output':str(target),'message':'관계·수량·단위 검증 통과'}
